from functools import reduce
from pathlib import Path
from openpyxl import load_workbook

CARPETA = Path(__file__).parent
ARCHIVO = CARPETA / "Dataset.xlsx"
SALIDA = CARPETA / "resultados.txt"

# Columnas del archivo
COL_NOMBRE = 0
COL_GRASAS = 3
COL_HIERRO = 5
COL_VIT_A = 6
COL_TIAMINA = 7
COL_RIBOFLAVINA = 8
COL_NIACINA = 9
COL_VIT_C = 10
COLS_VITAMINAS = [COL_VIT_A, COL_TIAMINA, COL_RIBOFLAVINA, COL_NIACINA, COL_VIT_C]

def a_numero(valor):
    """Convierte una celda a float; '-' o vacío se toma como 0."""
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0

def leer_filas(ruta):
    hoja = load_workbook(ruta, read_only=True, data_only=True).active
    filas = list(hoja.iter_rows(values_only=True))[1:]  # se omite el encabezado
    filas = filter(lambda f: f[COL_NOMBRE], filas)       # se omiten filas vacías
    return list(map(lambda f: {
        "nombre": str(f[COL_NOMBRE]).strip(),
        "grasas": a_numero(f[COL_GRASAS]),
        "hierro": a_numero(f[COL_HIERRO]),
        "tiamina": a_numero(f[COL_TIAMINA]),
        "vit_c": a_numero(f[COL_VIT_C]),
        "suma_vitaminas": sum(map(lambda c: a_numero(f[c]), COLS_VITAMINAS)),
    }, filas))

def maximo_por(datos, campo):
    pares = map(lambda d: ([d["nombre"]], d[campo]), datos)

    def mayor(a, b):
        if a[1] > b[1]:
            return a
        if b[1] > a[1]:
            return b
        return (a[0] + b[0], a[1])  # empate

    nombres, valor = reduce(mayor, pares)
    return " y ".join(nombres), valor

def main():
    datos = leer_filas(ARCHIVO)
    lineas = []  # aquí se acumula todo lo que irá al .txt

    lineas.append("RESULTADOS - ANÁLISIS CON MAP-REDUCE")
    lineas.append(f"Alimentos analizados: {len(datos)}\n")

    # 1. Alimento con mayor cantidad de grasas
    nombre, valor = maximo_por(datos, "grasas")
    lineas.append(f"1. Mayor cantidad de grasas: {nombre} ({valor:g} g)")

    # 2. Alimento con mayor cantidad de vitamina C
    nombre, valor = maximo_por(datos, "vit_c")
    lineas.append(f"2. Mayor cantidad de vitamina C: {nombre} ({valor:g} mg)")

    # 3. Cuántos alimentos tienen más de 0.1 de tiamina
    conteo = reduce(lambda acc, x: acc + x,
                    map(lambda d: 1 if d["tiamina"] > 0.1 else 0, datos), 0)
    lineas.append(f"3. Alimentos con tiamina > 0.1: {conteo}")

    # 4. Alimento con mayor suma de vitaminas (A + B1 + B2 + Niacina + C)
    nombre, valor = maximo_por(datos, "suma_vitaminas")
    lineas.append(f"4. Mayor suma de vitaminas: {nombre} ({valor:g})")

    # 5. Alimentos con más de 1 mg de hierro y menos de 3 g de grasas
    seleccion = reduce(
        lambda acc, par: acc + [par[0]] if par[1] else acc,
        map(lambda d: (d["nombre"], d["hierro"] > 1 and d["grasas"] < 3), datos),
        [],
    )
    lineas.append(f"5. Hierro > 1 mg y grasas < 3 g ({len(seleccion)} alimentos):")
    lineas.extend(map(lambda n: f"   - {n}", seleccion))

    SALIDA.write_text("\n".join(lineas) + "\n", encoding="utf-8")

if __name__ == "__main__":
    main()